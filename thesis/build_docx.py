"""
build_docx.py
--------------
Converts every thesis/chapter-*.md file (sorted by filename) into a single
formatted Word document, thesis/UniFlow-Thesis.docx.

Usage:
    pip install python-docx
    python thesis/build_docx.py

Supports the Markdown subset actually used in the chapters: #/##/### headings,
**bold**/*italic*, "- " bullet lists, "1. " numbered lists, pipe tables,
fenced ``` code blocks (used for the ASCII diagram), and "[1] ..." IEEE-style
reference entries. Anything else is emitted as a plain paragraph.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.shared import Inches, Pt
from docx.oxml.ns import qn

THESIS_DIR = Path(__file__).parent
OUTPUT_PATH = THESIS_DIR / "UniFlow-Thesis.docx"

TABLE_ROW_RE = re.compile(r"^\s*\|")
TABLE_SEP_RE = re.compile(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)+\|?\s*$")
BULLET_RE = re.compile(r"^-\s+(.*)")
NUMBERED_RE = re.compile(r"^\d+\.\s+(.*)")
REFERENCE_RE = re.compile(r"^\[(\d+)\]\s+(.*)")
HEADING_RE = re.compile(r"^(#{1,3})\s+(.*)")
INLINE_RE = re.compile(r"\*\*(.+?)\*\*|\*(.+?)\*|([^*]+)")


def add_formatted_runs(paragraph, text: str) -> None:
    for match in INLINE_RE.finditer(text):
        bold, italic, plain = match.groups()
        if bold is not None:
            paragraph.add_run(bold).bold = True
        elif italic is not None:
            paragraph.add_run(italic).italic = True
        elif plain:
            paragraph.add_run(plain)


def add_code_block(document: Document, lines: list[str]) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(0)
    for i, line in enumerate(lines):
        if i > 0:
            paragraph.add_run().add_break()
        run = paragraph.add_run(line if line else " ")
        run.font.name = "Consolas"
        run.font.size = Pt(9)


def add_table(document: Document, rows: list[list[str]]) -> None:
    n_cols = max(len(r) for r in rows)
    table = document.add_table(rows=len(rows), cols=n_cols)
    table.style = "Table Grid"
    for r, row_cells in enumerate(rows):
        for c in range(n_cols):
            cell_text = row_cells[c] if c < len(row_cells) else ""
            cell_paragraph = table.cell(r, c).paragraphs[0]
            add_formatted_runs(cell_paragraph, cell_text)
            if r == 0:
                for run in cell_paragraph.runs:
                    run.bold = True
    document.add_paragraph()


def parse_table_row(line: str) -> list[str]:
    stripped = line.strip().strip("|")
    return [cell.strip() for cell in stripped.split("|")]


def convert_chapter(document: Document, md_path: Path, *, first: bool) -> None:
    if not first:
        document.add_page_break()

    lines = md_path.read_text(encoding="utf-8").splitlines()
    i = 0
    paragraph_buffer: list[str] = []
    list_buffer: list[str] | None = None
    list_kind: str | None = None  # "bullet" | "numbered" | "reference"

    def flush_paragraph():
        if paragraph_buffer:
            text = " ".join(paragraph_buffer).strip()
            paragraph_buffer.clear()
            if text:
                p = document.add_paragraph()
                add_formatted_runs(p, text)

    def flush_list():
        nonlocal list_buffer, list_kind
        if list_buffer:
            style = {
                "bullet": "List Bullet",
                "numbered": "List Number",
                "reference": None,
            }[list_kind]
            for item in list_buffer:
                p = document.add_paragraph(style=style)
                if list_kind == "reference":
                    p.paragraph_format.left_indent = Inches(0.4)
                    p.paragraph_format.first_line_indent = Inches(-0.4)
                add_formatted_runs(p, item)
        list_buffer = None
        list_kind = None

    while i < len(lines):
        line = lines[i]

        if line.strip() == "":
            flush_paragraph()
            flush_list()
            i += 1
            continue

        heading_match = HEADING_RE.match(line)
        if heading_match:
            flush_paragraph()
            flush_list()
            level = len(heading_match.group(1))
            document.add_heading(heading_match.group(2).strip(), level=level)
            i += 1
            continue

        if line.strip().startswith("```"):
            flush_paragraph()
            flush_list()
            i += 1
            code_lines = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            i += 1  # skip closing fence
            add_code_block(document, code_lines)
            continue

        if TABLE_ROW_RE.match(line):
            flush_paragraph()
            flush_list()
            table_lines = []
            while i < len(lines) and TABLE_ROW_RE.match(lines[i]):
                table_lines.append(lines[i])
                i += 1
            rows = [parse_table_row(l) for l in table_lines if not TABLE_SEP_RE.match(l)]
            add_table(document, rows)
            continue

        ref_match = REFERENCE_RE.match(line)
        bullet_match = BULLET_RE.match(line)
        numbered_match = NUMBERED_RE.match(line)

        if ref_match or bullet_match or numbered_match:
            flush_paragraph()
            kind = "reference" if ref_match else ("bullet" if bullet_match else "numbered")
            content = (ref_match or bullet_match or numbered_match).group(0) if kind == "reference" else (ref_match or bullet_match or numbered_match).group(1)
            if kind == "reference":
                content = line.strip()
            if list_kind != kind:
                flush_list()
                list_kind = kind
                list_buffer = []
            list_buffer.append(content)
            i += 1
            continue

        if list_buffer is not None:
            # Continuation line of the current (indented, wrapped) list item.
            list_buffer[-1] = f"{list_buffer[-1]} {line.strip()}"
            i += 1
            continue

        paragraph_buffer.append(line.strip())
        i += 1

    flush_paragraph()
    flush_list()


def main() -> None:
    chapter_files = sorted(THESIS_DIR.glob("chapter-*.md"))
    if not chapter_files:
        raise SystemExit(f"No chapter-*.md files found in {THESIS_DIR}")

    document = Document()

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.space_after = Pt(6)
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), "Times New Roman")

    for level, size in ((1, 18), (2, 14), (3, 12)):
        heading_style = document.styles[f"Heading {level}"]
        heading_style.font.name = "Times New Roman"
        heading_style.font.size = Pt(size)
        heading_style.font.bold = True
        heading_style.font.italic = level == 3
        heading_style.font.color.rgb = None

    for i, path in enumerate(chapter_files):
        convert_chapter(document, path, first=(i == 0))
        print(f"Added {path.name}")

    document.save(OUTPUT_PATH)
    print(f"\nWrote {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
